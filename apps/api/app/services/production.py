from __future__ import annotations

import html
import io
import ipaddress
import socket
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import urlparse
from urllib.request import Request, urlopen
from uuid import uuid4

from app.core.config import get_settings

import polars as pl

from app.services.profiling import profile_dataset
from app.db.models import AuditLog, DataSource, Dataset, DatasetVersion, Workspace
from sqlalchemy import func, select

from app.services.scheduler_lock import REFRESH_JOB_LOCK, release_job_lock, try_job_lock

settings = get_settings()


def validate_public_url(url: str) -> None:
    parsed = urlparse(url)
    if parsed.scheme not in {"http", "https"} or not parsed.hostname:
        raise ValueError("Connector URL must use http or https")
    host = parsed.hostname
    if host.lower() in {"localhost", "localhost.localdomain"}:
        raise ValueError("Private/local connector hosts are not allowed")
    try:
        infos = socket.getaddrinfo(host, parsed.port or (443 if parsed.scheme == "https" else 80), type=socket.SOCK_STREAM)
    except OSError as exc:
        raise ValueError(f"Could not resolve connector host: {exc}") from exc
    for info in infos:
        addr = ipaddress.ip_address(info[4][0])
        if addr.is_private or addr.is_loopback or addr.is_link_local or addr.is_reserved or addr.is_multicast:
            raise ValueError("Connector URL resolves to a private or reserved network address")


def download_csv(url: str, destination: Path, max_bytes: int) -> int:
    validate_public_url(url)
    req = Request(url, headers={"User-Agent": "AI-Analytics-Platform/1.0"})
    total = 0
    with urlopen(req, timeout=30) as response, destination.open("wb") as output:
        final_url = response.geturl()
        if final_url != url:
            validate_public_url(final_url)
        while chunk := response.read(1024 * 1024):
            total += len(chunk)
            if total > max_bytes:
                raise ValueError("Remote file exceeds the configured upload limit")
            output.write(chunk)
    return total


def export_csv(frame: pl.DataFrame) -> bytes:
    buf = io.StringIO()
    frame.write_csv(buf)
    return buf.getvalue().encode("utf-8")


def export_json(frame: pl.DataFrame) -> bytes:
    return frame.write_json().encode("utf-8")


def export_xlsx(frame: pl.DataFrame) -> bytes:
    # Polars delegates Excel writing to xlsxwriter; keep it isolated to exports.
    import xlsxwriter

    buf = io.BytesIO()
    with xlsxwriter.Workbook(buf, {"in_memory": True}) as workbook:
        sheet = workbook.add_worksheet("Data")
        for c, name in enumerate(frame.columns):
            sheet.write(0, c, name)
        for r, row in enumerate(frame.iter_rows(), start=1):
            for c, value in enumerate(row):
                if hasattr(value, "isoformat"):
                    value = value.isoformat()
                sheet.write(r, c, value)
        sheet.freeze_panes(1, 0)
        sheet.autofilter(0, 0, max(0, frame.height), max(0, frame.width - 1))
    return buf.getvalue()


def write_report_html(dataset_name: str, frame: pl.DataFrame, profile: dict) -> bytes:
    missing = sum(int(x.get("count", 0)) for x in profile.get("missing", []))
    if frame.width:
        rows = frame.head(25).rows()
        th = "".join(f"<th>{html.escape(str(name))}</th>" for name in frame.columns)
        body = "".join("<tr>" + "".join(f"<td>{html.escape('' if value is None else str(value))}</td>" for value in row) + "</tr>" for row in rows)
        head = f"<table class='data'><thead><tr>{th}</tr></thead><tbody>{body}</tbody></table>"
    else:
        head = "<p>No columns.</p>"
    html = f"""<!doctype html><html><head><meta charset='utf-8'><title>{html.escape(str(dataset_name))} report</title>
    <style>body{{font-family:Arial,sans-serif;margin:40px;color:#172033}}.metrics{{display:flex;gap:12px;flex-wrap:wrap}}.metric{{border:1px solid #ddd;border-radius:10px;padding:14px;min-width:120px}}.data{{border-collapse:collapse;width:100%;font-size:12px}}.data th,.data td{{border:1px solid #ddd;padding:6px;text-align:left}}.data th{{background:#f4f6f8}}</style></head>
    <body><h1>{html.escape(str(dataset_name))}</h1><p>Generated {datetime.now(timezone.utc).isoformat()}</p>
    <div class='metrics'><div class='metric'><b>Rows</b><br>{profile.get('rows',0):,}</div><div class='metric'><b>Columns</b><br>{profile.get('columns',0)}</div><div class='metric'><b>Missing cells</b><br>{missing:,}</div><div class='metric'><b>Quality</b><br>{profile.get('quality_score',0)}%</div></div>
    <h2>Preview</h2>{head}</body></html>"""
    return html.encode("utf-8")



def refresh_source(db, source: DataSource, actor=None) -> Dataset:
    workspace = db.get(Workspace, source.workspace_id)
    if not workspace:
        raise ValueError("Workspace not found")
    dataset = db.get(Dataset, source.dataset_id) if source.dataset_id else None
    if dataset is None:
        destination = settings.upload_path / f"connector-{uuid4()}.csv"
        try:
            size = download_csv(source.url, destination, settings.max_upload_mb * 1024 * 1024)
            profile = profile_dataset(destination)
            dataset = Dataset(
                workspace_id=workspace.id, name=source.name,
                filename=Path(urlparse(source.url).path).name or f"{source.name}.csv",
                storage_path=str(destination), file_type="csv", file_size_bytes=size,
                row_count=profile["rows"], column_count=profile["columns"], profile=profile,
            )
            db.add(dataset)
            db.flush()
            version_no = 1
        except Exception:
            destination.unlink(missing_ok=True)
            raise
    else:
        destination = settings.processed_path / str(dataset.id) / f"connector-{uuid4()}.csv"
        destination.parent.mkdir(parents=True, exist_ok=True)
        try:
            size = download_csv(source.url, destination, settings.max_upload_mb * 1024 * 1024)
            profile = profile_dataset(destination)
        except Exception:
            destination.unlink(missing_ok=True)
            raise
        version_no = (db.scalar(select(func.max(DatasetVersion.version)).where(DatasetVersion.dataset_id == dataset.id)) or 0) + 1
        dataset.storage_path = str(destination)
        dataset.file_type = "csv"
        dataset.file_size_bytes = size
        dataset.row_count = profile["rows"]
        dataset.column_count = profile["columns"]
        dataset.profile = profile
    db.add(DatasetVersion(
        dataset_id=dataset.id, version=version_no, label=f"Connector refresh {version_no}",
        storage_path=str(destination), file_type="csv", file_size_bytes=size,
        row_count=profile["rows"], column_count=profile["columns"], profile=profile,
    ))
    source.dataset_id = dataset.id
    db.add(AuditLog(
        user_id=actor.id if actor else None, action="connector.refresh",
        resource_type="dataset", resource_id=str(dataset.id),
        details={"source_id": str(source.id), "version": version_no},
    ))
    return dataset

def run_due_refresh_jobs() -> None:
    from datetime import timedelta
    from sqlalchemy import select
    from app.db.database import SessionLocal
    from app.db.models import RefreshJob
    db = SessionLocal()
    locked = False
    try:
        locked = try_job_lock(db, REFRESH_JOB_LOCK)
        if not locked:
            return
        now = datetime.now(timezone.utc)
        jobs = db.scalars(select(RefreshJob).where(RefreshJob.enabled.is_(True), RefreshJob.next_run_at <= now).limit(20)).all()
        for job in jobs:
            source = db.get(DataSource, job.source_id)
            if not source or not source.enabled:
                job.last_status = "skipped"
                job.last_error = "Source is missing or disabled"
            else:
                try:
                    refresh_source(db, source, None)
                    job.last_status = "success"
                    job.last_error = None
                except Exception as exc:
                    job.last_status = "error"
                    job.last_error = str(exc)[:1000]
            job.last_run_at = now
            job.next_run_at = now + timedelta(minutes=job.interval_minutes)
        db.commit()
    finally:
        db.close()


def export_pdf(dataset_name: str, profile: dict) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    width, height = A4
    c.setTitle(f"{dataset_name} report")
    c.setFont("Helvetica-Bold", 20)
    c.drawString(48, height - 60, dataset_name)
    c.setFont("Helvetica", 11)
    y = height - 95
    metrics = [("Rows", profile.get("rows", 0)), ("Columns", profile.get("columns", 0)),
               ("Quality", f"{profile.get('quality_score', 0)}%"), ("Missing", f"{profile.get('missing_percent', 0)}%")] 
    for label, value in metrics:
        c.drawString(55, y, f"{label}: {value}")
        y -= 24
    c.drawString(55, y - 10, f"Generated: {datetime.now(timezone.utc).isoformat()}")
    c.save()
    return buf.getvalue()


def export_pptx(dataset_name: str, profile: dict) -> bytes:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = dataset_name
    box = slide.shapes.add_textbox(Inches(1), Inches(1.6), Inches(8), Inches(3.5))
    tf = box.text_frame
    tf.clear()
    for label, value in [("Rows", profile.get("rows", 0)), ("Columns", profile.get("columns", 0)),
                         ("Quality", f"{profile.get('quality_score', 0)}%"), ("Missing", f"{profile.get('missing_percent', 0)}%")]:
        p = tf.add_paragraph() if tf.text else tf.paragraphs[0]
        p.text = f"{label}: {value}"
        p.font.size = Pt(22)
    buf = io.BytesIO(); prs.save(buf); return buf.getvalue()
